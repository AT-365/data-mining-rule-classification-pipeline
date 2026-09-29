from __future__ import annotations

import copy
import json
import math
import sys
from collections import Counter
from datetime import datetime
from html import escape
from pathlib import Path
from typing import TYPE_CHECKING, Any, Dict, List, Tuple

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import openpyxl  # noqa: F401
except ImportError:
    openpyxl = None

try:
    import docx  # noqa: F401
except ImportError:
    docx = None

try:
    import reportlab  # noqa: F401
except ImportError:
    reportlab = None

if TYPE_CHECKING:
    from pandas import DataFrame, Series


CONFIG = {
    "base_dir": Path(__file__).resolve().parent,
    "data_dir": "data",
    "outputs_dir": "outputs",
    "reports_dir": "reports",
    "private_notes_dir": "defense_notes",
    "id3_train_csv": "A2--ID3-Training set.csv",
    "id3_test_csv": "A2--ID3-TEST set.csv",
    "bayes_train_csv": "A2--Bayes-Training set.csv",
    "bayes_test_csv": "A2--Bayes-TEST set.csv",
    "class_attribute": "Volume",
    "bayes_continuous_attributes": ["Weight", "Material"],
    "show_live_section_output_default": True,
    "pdf_line_width": 96,
}

SESSION_FLAGS = {
    "show_live_terminal_output": CONFIG["show_live_section_output_default"],
}


def validate_runtime_environment() -> None:
    """Stop early with a clear message when Python or packages are missing."""
    if sys.version_info < (3, 10):
        raise RuntimeError(
            "Python 3.10 or newer is required. "
            f"Current version: {sys.version.split()[0]}"
        )

    missing_packages = []
    if pd is None:
        missing_packages.append("pandas")
    if openpyxl is None:
        missing_packages.append("openpyxl")
    if docx is None:
        missing_packages.append("python-docx")
    if reportlab is None:
        missing_packages.append("reportlab")

    if missing_packages:
        raise ImportError(
            "Missing required packages: "
            + ", ".join(missing_packages)
            + ". Install them with: python -m pip install -r requirements.txt"
        )


def validate_project_files(base_dir: Path) -> None:
    """Confirm the expected assignment data files are present before prompting."""
    data_dir = base_dir / CONFIG["data_dir"]
    required_files = [
        CONFIG["id3_train_csv"],
        CONFIG["id3_test_csv"],
        CONFIG["bayes_train_csv"],
        CONFIG["bayes_test_csv"],
    ]

    missing_files = [str(data_dir / file_name) for file_name in required_files if not (data_dir / file_name).exists()]
    if missing_files:
        raise FileNotFoundError(
            "Missing required data files: " + ", ".join(missing_files)
        )


def explain_function_purpose() -> Dict[str, str]:
    """Return a plain-language map of the major functions used in this pipeline."""
    return {
        "get_classifier_choice": "Asks the user whether to run ID3 or Naive Bayes before any calculations begin.",
        "get_decimal_input": "Collects decimal threshold inputs and checks that the value is inside the required range.",
        "load_selected_datasets": "Loads the correct training and test files based on the selected classifier.",
        "inspect_dataset": "Summarizes row counts, columns, missing values, and Volume class counts for proof that the right files loaded.",
        "calculate_entropy": "Calculates the entropy or MC value used by ID3 to measure class mixture.",
        "choose_best_id3_attribute": "Calculates WMC and Gain for candidate attributes and selects the highest-gain split.",
        "calculate_mixture_ratio": "Calculates alpha = c1 / |F1| for ID3 pre-pruning decisions.",
        "build_id3_tree": "Builds the ID3 decision tree from the discretized ID3 training set using pre-pruning.",
        "extract_id3_rules": "Converts the final decision tree into readable IF-THEN rules.",
        "predict_with_id3_tree": "Classifies test records by walking through the ID3 tree from root to leaf.",
        "apply_post_pruning": "Tests whether subtrees can be removed using the assignment rule (N-M)/Q <= gK.",
        "identify_bayes_attribute_types": "Separates Naive Bayes attributes into categorical attributes and continuous attributes.",
        "calculate_class_priors": "Calculates the starting chance of each Volume class from the training set.",
        "check_smoothing_needed": "Checks whether a zero categorical probability would occur without smoothing.",
        "calculate_categorical_probabilities": "Calculates categorical probability tables from counts, using Laplace smoothing only when needed.",
        "calculate_continuous_statistics": "Calculates mean and sample standard deviation by Volume class for continuous attributes.",
        "calculate_continuous_probability": "Calculates the Gaussian probability component for continuous values using the standard math library.",
        "predict_bayes_test_records": "Calculates one Naive Bayes score per Volume class and predicts the class with the highest score.",
        "calculate_metrics": "Builds the multiclass confusion matrix and accuracy tables for the test predictions.",
        "save_excel_output": "Saves each major calculation step as its own Excel workbook and prints the saved path.",
        "save_report_files": "Creates the TXT, MD, HTML, DOCX, and PDF report files for the selected method.",
        "build_defense_notes": "Creates a quick-glance Zoom guide explaining who, what, when, where, why, and how for the selected method.",
    }


def live_print(enabled: bool, message: str = "") -> None:
    """Print a message only when live terminal output is turned on."""
    if enabled:
        print(message, flush=True)


def should_print_saved_file(file_path: Path) -> bool:
    """Return whether saved-file messages should be printed to the terminal."""
    return SESSION_FLAGS["show_live_terminal_output"]


def print_saved_file(file_path: Path) -> None:
    """Print the saved file path immediately so the VS Code terminal shows progress."""
    if not should_print_saved_file(file_path):
        return
    try:
        display_path = file_path.resolve().relative_to(CONFIG["base_dir"])
    except ValueError:
        display_path = file_path
    print(f"Saved: {display_path}", flush=True)


def get_yes_no_input(prompt_text: str, default_value: bool = True) -> bool:
    """Ask a yes/no question and return a boolean value."""
    default_text = "yes" if default_value else "no"
    while True:
        user_input = input(f"{prompt_text} (yes/no, default {default_text}): ").strip().lower()
        if not user_input:
            return default_value
        if user_input in {"yes", "y", "true", "t"}:
            return True
        if user_input in {"no", "n", "false", "f"}:
            return False
        print("Invalid entry. Please enter yes or no.\n", flush=True)


def get_classifier_choice() -> str:
    """Ask the user which classifier to run and normalize the answer."""
    print("--- Assignment 2 Initialization ---")
    print("Choose a classifier: ID3 or Naive Bayes.")
    print("ID3 builds a decision tree and readable IF-THEN rules.")
    print("Naive Bayes calculates probability scores and predicts the most likely Volume class.\n")

    while True:
        user_input = input("Enter classifier choice (ID3 or Bayes): ").strip().lower()
        if user_input in {"id3", "decision tree", "tree"}:
            return "id3"
        if user_input in {"bayes", "naive bayes", "naïve bayes", "naives bayes", "naive", "nb"}:
            return "bayes"
        print("Invalid entry. Please enter ID3 or Bayes.\n", flush=True)


def get_decimal_input(prompt_text: str, lower_bound: float, upper_bound: float, include_upper: bool = True) -> float:
    """Collect a decimal input and validate that it is inside the requested range."""
    while True:
        user_input = input(prompt_text).strip()
        try:
            value = float(user_input)
        except ValueError:
            print("Invalid entry. Please enter a decimal number.\n", flush=True)
            continue

        lower_ok = value > lower_bound
        upper_ok = value <= upper_bound if include_upper else value < upper_bound
        if lower_ok and upper_ok:
            return value

        upper_symbol = "<=" if include_upper else "<"
        print(f"Invalid entry. Valid range: {lower_bound} < value {upper_symbol} {upper_bound}.\n", flush=True)


def collect_startup_inputs() -> Tuple[str, bool, Dict[str, Any]]:
    """Collect every required user input before the pipeline starts processing files."""
    classifier_choice = get_classifier_choice()
    live_output = True
    print("Live terminal progress output is enabled automatically.", flush=True)

    user_inputs: Dict[str, Any] = {
        "Classifier Choice": "ID3" if classifier_choice == "id3" else "Naive Bayes",
        "Live Terminal Output": "Yes" if live_output else "No",
    }

    if classifier_choice == "id3":
        t1_threshold = get_decimal_input(
            "Enter the pre-pruning threshold T1. T1 controls when ID3 stops splitting early using alpha = c1 / |F1|. Enter T1 (0 < T1 <= 1): ",
            0,
            1,
        )
        g_parameter = get_decimal_input(
            "Enter the post-pruning parameter g. Valid range: 0 < g <= 0.015. Enter g: ",
            0,
            0.015,
        )
        user_inputs["T1 Threshold"] = t1_threshold
        user_inputs["g Parameter"] = g_parameter

    return classifier_choice, live_output, user_inputs


def create_assignment_folders(base_dir: Path) -> Dict[str, Path]:
    """Create the folders used by the VS Code project and return their paths."""
    data_dir = base_dir / CONFIG["data_dir"]
    outputs_dir = base_dir / CONFIG["outputs_dir"]
    reports_dir = base_dir / CONFIG["reports_dir"]
    private_notes_dir = base_dir / CONFIG["private_notes_dir"]

    for folder in [
        data_dir,
        outputs_dir,
        reports_dir,
        private_notes_dir,
        outputs_dir / "id3",
        outputs_dir / "bayes",
        reports_dir / "id3",
        reports_dir / "bayes",
        private_notes_dir / "id3",
        private_notes_dir / "bayes",
    ]:
        folder.mkdir(parents=True, exist_ok=True)

    return {
        "data": data_dir,
        "outputs": outputs_dir,
        "reports": reports_dir,
        "private_notes": private_notes_dir,
        "id3_outputs": outputs_dir / "id3",
        "bayes_outputs": outputs_dir / "bayes",
        "id3_reports": reports_dir / "id3",
        "bayes_reports": reports_dir / "bayes",
        "id3_private_notes": private_notes_dir / "id3",
        "bayes_private_notes": private_notes_dir / "bayes",
    }


def print_dataframe_output(title: str, dataframe: DataFrame) -> None:
    """Print one DataFrame to the terminal so results are visible without opening Excel."""
    if not SESSION_FLAGS["show_live_terminal_output"]:
        return
    print(f"\n--- {title} ---", flush=True)
    if dataframe.empty:
        print("[No rows]", flush=True)
        return
    with pd.option_context(
        "display.max_rows", None,
        "display.max_columns", None,
        "display.width", None,
        "display.max_colwidth", None,
    ):
        print(dataframe.to_string(index=False), flush=True)


def save_excel_output(file_path: Path, payload: DataFrame | Dict[str, DataFrame]) -> None:
    """Save one DataFrame or several named sheets to an Excel workbook."""
    if isinstance(payload, dict):
        with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
            for sheet_name, dataframe in payload.items():
                safe_sheet_name = sheet_name[:31]
                dataframe.to_excel(writer, sheet_name=safe_sheet_name, index=False)
    else:
        payload.to_excel(file_path, index=False)
    print_saved_file(file_path)
    if isinstance(payload, dict):
        for sheet_name, dataframe in payload.items():
            print_dataframe_output(f"{file_path.name} [{sheet_name}]", dataframe)
    else:
        print_dataframe_output(file_path.name, payload)


def save_text_output(file_path: Path, text: str) -> None:
    """Save text output and print the file path after it is written."""
    file_path.write_text(text, encoding="utf-8")
    print_saved_file(file_path)


def load_selected_datasets(classifier_choice: str, folders: Dict[str, Path]) -> Tuple[DataFrame, DataFrame, str, str]:
    """Load the matching training and test files for ID3 or Naive Bayes."""
    data_dir = folders["data"]
    if classifier_choice == "id3":
        train_file = CONFIG["id3_train_csv"]
        test_file = CONFIG["id3_test_csv"]
    else:
        train_file = CONFIG["bayes_train_csv"]
        test_file = CONFIG["bayes_test_csv"]

    train_path = data_dir / train_file
    test_path = data_dir / test_file
    if not train_path.exists():
        raise FileNotFoundError(f"Missing training file: {train_path}")
    if not test_path.exists():
        raise FileNotFoundError(f"Missing test file: {test_path}")

    training_set = pd.read_csv(train_path)
    test_set = pd.read_csv(test_path)
    return training_set, test_set, train_file, test_file


def inspect_dataset(dataframe: DataFrame, dataset_name: str) -> Dict[str, DataFrame]:
    """Build compact audit tables for a dataset."""
    class_attribute = CONFIG["class_attribute"]
    overview_df = pd.DataFrame([
        {"Dataset": dataset_name, "Row Count": len(dataframe), "Column Count": len(dataframe.columns)},
    ])
    columns_df = pd.DataFrame({
        "Column Name": dataframe.columns.tolist(),
        "Data Type": [str(dataframe[column].dtype) for column in dataframe.columns],
        "Missing Values": [int(dataframe[column].isna().sum()) for column in dataframe.columns],
        "Unique Values": [int(dataframe[column].nunique()) for column in dataframe.columns],
    })
    if class_attribute in dataframe.columns:
        class_counts_df = dataframe[class_attribute].value_counts().sort_index().reset_index()
        class_counts_df.columns = [class_attribute, "Record Count"]
    else:
        class_counts_df = pd.DataFrame(columns=[class_attribute, "Record Count"])
    return {"Overview": overview_df, "Columns": columns_df, "Volume Counts": class_counts_df}


def calculate_class_distribution(dataframe: DataFrame, class_attribute: str) -> DataFrame:
    """Count the number and percentage of records in each class value."""
    total_records = len(dataframe)
    rows = []
    for class_value, count in sorted(Counter(dataframe[class_attribute]).items()):
        rows.append({
            "Class Attribute": class_attribute,
            "Class Value": class_value,
            "Record Count": count,
            "Percent of Records": count / total_records if total_records else 0,
        })
    return pd.DataFrame(rows)


def calculate_entropy(class_values: List[Any]) -> float:
    """Calculate entropy/MC for a list of class labels."""
    if not class_values:
        return 0.0
    total_count = len(class_values)
    class_counts = Counter(class_values)
    entropy_value = 0.0
    for count in class_counts.values():
        probability = count / total_count
        entropy_value -= probability * math.log2(probability)
    return entropy_value


def format_class_counts(class_values: List[Any]) -> str:
    """Format class counts so they are readable in output tables."""
    counts = Counter(class_values)
    return ", ".join(f"{class_value}:{counts[class_value]}" for class_value in sorted(counts))


def dominant_class_and_count(class_values: List[Any]) -> Tuple[Any, int]:
    """Return the dominant class and its count, using the smaller class value to break ties."""
    counts = Counter(class_values)
    if not counts:
        return None, 0
    sorted_counts = sorted(counts.items(), key=lambda pair: (-pair[1], pair[0]))
    return sorted_counts[0]


def calculate_weighted_entropy(dataframe: DataFrame, split_attribute: str, class_attribute: str) -> Tuple[float, str]:
    """Calculate weighted entropy for a candidate ID3 split attribute."""
    total_count = len(dataframe)
    weighted_entropy = 0.0
    partition_pieces = []
    for attribute_value, partition_df in sorted(dataframe.groupby(split_attribute), key=lambda pair: pair[0]):
        partition_entropy = calculate_entropy(partition_df[class_attribute].tolist())
        weight = len(partition_df) / total_count if total_count else 0
        weighted_entropy += weight * partition_entropy
        partition_pieces.append(
            f"{split_attribute}={attribute_value}: n={len(partition_df)}, MC={partition_entropy:.6f}"
        )
    return weighted_entropy, " | ".join(partition_pieces)


def choose_best_id3_attribute(
    dataframe: DataFrame,
    candidate_attributes: List[str],
    class_attribute: str,
    node_id: int,
) -> Tuple[str | None, List[Dict[str, Any]]]:
    """Calculate gain for candidate attributes and return the highest-gain attribute."""
    parent_entropy = calculate_entropy(dataframe[class_attribute].tolist())
    gain_rows = []
    best_attribute = None
    best_gain = -1.0

    for attribute in candidate_attributes:
        weighted_entropy, partition_details = calculate_weighted_entropy(dataframe, attribute, class_attribute)
        information_gain = parent_entropy - weighted_entropy
        gain_rows.append({
            "Node ID": node_id,
            "Candidate Attribute": attribute,
            "Parent MC": parent_entropy,
            "WMC(Attribute)": weighted_entropy,
            "Gain": information_gain,
            "Partition Details": partition_details,
        })
        if information_gain > best_gain:
            best_gain = information_gain
            best_attribute = attribute

    for row in gain_rows:
        row["Selected Attribute"] = "Yes" if row["Candidate Attribute"] == best_attribute else "No"
    return best_attribute, gain_rows


def calculate_mixture_ratio(dataframe: DataFrame, class_attribute: str) -> Tuple[Any, int, float]:
    """Calculate alpha = c1 / |F1| for a possible ID3 branch."""
    class_values = dataframe[class_attribute].tolist()
    dominant_class, dominant_count = dominant_class_and_count(class_values)
    alpha = dominant_count / len(dataframe) if len(dataframe) else 0.0
    return dominant_class, dominant_count, alpha


def build_leaf_node(
    node_id: int,
    depth: int,
    dataframe: DataFrame,
    class_attribute: str,
    predicted_class: Any,
    reason: str,
    path_text: str,
) -> Dict[str, Any]:
    """Build a leaf node dictionary with enough detail for reports and predictions."""
    return {
        "node_id": node_id,
        "type": "leaf",
        "depth": depth,
        "attribute": None,
        "children": {},
        "predicted_class": predicted_class,
        "dominant_class": predicted_class,
        "record_count": len(dataframe),
        "class_counts": dict(Counter(dataframe[class_attribute].tolist())) if len(dataframe) else {},
        "reason": reason,
        "path_text": path_text,
    }


def build_id3_tree(
    dataframe: DataFrame,
    candidate_attributes: List[str],
    class_attribute: str,
    t1_threshold: float,
    node_counter: List[int],
    gain_rows: List[Dict[str, Any]],
    split_rows: List[Dict[str, Any]],
    pre_pruning_rows: List[Dict[str, Any]],
    depth: int = 0,
    path_text: str = "ROOT",
) -> Dict[str, Any]:
    """Build an ID3 tree from scratch using the assignment's pre-pruning rule."""
    node_counter[0] += 1
    node_id = node_counter[0]
    class_values = dataframe[class_attribute].tolist()
    dominant_class, dominant_count = dominant_class_and_count(class_values)

    if len(dataframe) == 0:
        return build_leaf_node(node_id, depth, dataframe, class_attribute, dominant_class, "Empty partition", path_text)

    if len(set(class_values)) == 1:
        return build_leaf_node(node_id, depth, dataframe, class_attribute, class_values[0], "Pure class leaf", path_text)

    if not candidate_attributes:
        return build_leaf_node(node_id, depth, dataframe, class_attribute, dominant_class, "No attributes left", path_text)

    selected_attribute, node_gain_rows = choose_best_id3_attribute(dataframe, candidate_attributes, class_attribute, node_id)
    gain_rows.extend(node_gain_rows)

    if selected_attribute is None:
        return build_leaf_node(node_id, depth, dataframe, class_attribute, dominant_class, "No selected attribute", path_text)

    selected_gain = max(row["Gain"] for row in node_gain_rows if row["Candidate Attribute"] == selected_attribute)
    if selected_gain <= 0:
        return build_leaf_node(node_id, depth, dataframe, class_attribute, dominant_class, "No positive information gain", path_text)

    node = {
        "node_id": node_id,
        "type": "decision",
        "depth": depth,
        "attribute": selected_attribute,
        "children": {},
        "predicted_class": None,
        "dominant_class": dominant_class,
        "record_count": len(dataframe),
        "class_counts": dict(Counter(class_values)),
        "reason": "Decision node selected by highest Gain",
        "path_text": path_text,
    }
    split_rows.append({
        "Node ID": node_id,
        "Depth": depth,
        "Path": path_text,
        "Selected Attribute": selected_attribute,
        "Gain": selected_gain,
        "Record Count": len(dataframe),
        "Dominant Volume": dominant_class,
        "Class Counts": format_class_counts(class_values),
    })

    remaining_attributes = [attribute for attribute in candidate_attributes if attribute != selected_attribute]
    for attribute_value, child_df in sorted(dataframe.groupby(selected_attribute), key=lambda pair: pair[0]):
        child_path = f"{path_text} AND {selected_attribute}={attribute_value}" if path_text != "ROOT" else f"{selected_attribute}={attribute_value}"
        branch_class, branch_dominant_count, alpha = calculate_mixture_ratio(child_df, class_attribute)
        should_pre_prune = alpha < t1_threshold
        pre_pruning_rows.append({
            "Parent Node ID": node_id,
            "Split Attribute": selected_attribute,
            "Branch Value": attribute_value,
            "Branch Path": child_path,
            "Branch Record Count |F1|": len(child_df),
            "Dominant Class c1": branch_class,
            "Dominant Class Count": branch_dominant_count,
            "Alpha = c1 / |F1|": alpha,
            "T1 Threshold": t1_threshold,
            "Pre-Pruned?": "Yes" if should_pre_prune else "No",
            "Decision": "Stop and assign dominant Volume" if should_pre_prune else "Continue if other stop conditions do not apply",
        })

        if should_pre_prune:
            node_counter[0] += 1
            child_node = build_leaf_node(
                node_counter[0],
                depth + 1,
                child_df,
                class_attribute,
                branch_class,
                "Pre-pruned because alpha < T1",
                child_path,
            )
        else:
            child_node = build_id3_tree(
                child_df,
                remaining_attributes,
                class_attribute,
                t1_threshold,
                node_counter,
                gain_rows,
                split_rows,
                pre_pruning_rows,
                depth + 1,
                child_path,
            )
        node["children"][attribute_value] = child_node

    return node


def flatten_tree(tree: Dict[str, Any]) -> DataFrame:
    """Convert the ID3 tree dictionary into a readable table."""
    rows = []

    def walk(node: Dict[str, Any], parent_node_id: Any = "", branch_value: Any = "") -> None:
        rows.append({
            "Node ID": node["node_id"],
            "Parent Node ID": parent_node_id,
            "Branch Value From Parent": branch_value,
            "Depth": node["depth"],
            "Node Type": node["type"],
            "Split Attribute": node.get("attribute", ""),
            "Predicted Volume": node.get("predicted_class", ""),
            "Dominant Volume": node.get("dominant_class", ""),
            "Record Count": node.get("record_count", 0),
            "Class Counts": node.get("class_counts", {}),
            "Path": node.get("path_text", ""),
            "Reason": node.get("reason", ""),
        })
        for child_value, child_node in node.get("children", {}).items():
            walk(child_node, node["node_id"], child_value)

    walk(tree)
    return pd.DataFrame(rows)


def extract_id3_rules(tree: Dict[str, Any]) -> DataFrame:
    """Convert every root-to-leaf path in the ID3 tree into an IF-THEN rule."""
    rows = []

    def walk(node: Dict[str, Any], conditions: List[str]) -> None:
        if node["type"] == "leaf":
            condition_text = " AND ".join(conditions) if conditions else "TRUE"
            rows.append({
                "Rule Number": len(rows) + 1,
                "Conditions": condition_text,
                "Predicted Volume": node["predicted_class"],
                "Record Count at Leaf": node["record_count"],
                "Class Counts at Leaf": node["class_counts"],
                "Rule Text": f"IF {condition_text} THEN Volume={node['predicted_class']}",
                "Leaf Node ID": node["node_id"],
                "Leaf Reason": node["reason"],
            })
            return

        split_attribute = node["attribute"]
        for child_value, child_node in sorted(node["children"].items(), key=lambda pair: pair[0]):
            walk(child_node, conditions + [f"{split_attribute}={child_value}"])

    walk(tree, [])
    return pd.DataFrame(rows)


def predict_one_id3_record(tree: Dict[str, Any], record: Series) -> Tuple[Any, str, str]:
    """Predict one record by walking through the ID3 tree."""
    node = tree
    path_steps = []
    while node["type"] == "decision":
        split_attribute = node["attribute"]
        record_value = record[split_attribute]
        path_steps.append(f"{split_attribute}={record_value}")
        if record_value not in node["children"]:
            return node["dominant_class"], "Missing branch, used node dominant class", " AND ".join(path_steps)
        node = node["children"][record_value]
    return node["predicted_class"], node["reason"], " AND ".join(path_steps)


def predict_with_id3_tree(tree: Dict[str, Any], test_set: DataFrame, class_attribute: str) -> DataFrame:
    """Predict every test record with the ID3 tree and keep the path used."""
    prediction_rows = []
    for row_number, (_, row) in enumerate(test_set.iterrows(), start=1):
        predicted_class, prediction_reason, path_used = predict_one_id3_record(tree, row)
        actual_class = row[class_attribute]
        prediction_rows.append({
            "Test Row Number": row_number,
            "Actual Volume": actual_class,
            "Predicted Volume": predicted_class,
            "Correct?": "Yes" if actual_class == predicted_class else "No",
            "Prediction Path": path_used,
            "Prediction Reason": prediction_reason,
            **{column: row[column] for column in test_set.columns},
        })
    return pd.DataFrame(prediction_rows)


def count_correct_predictions(prediction_df: DataFrame) -> int:
    """Count the number of correct predictions in a prediction table."""
    return int((prediction_df["Actual Volume"] == prediction_df["Predicted Volume"]).sum())


def count_subtree_branches(node: Dict[str, Any]) -> int:
    """Count child-branch edges inside a subtree for the post-pruning K value."""
    if node["type"] == "leaf":
        return 0
    branch_count = len(node["children"])
    for child_node in node["children"].values():
        branch_count += count_subtree_branches(child_node)
    return branch_count


def collect_decision_node_ids(tree: Dict[str, Any]) -> List[int]:
    """Collect decision node IDs from deepest to shallowest so post-pruning works bottom-up."""
    rows = []

    def walk(node: Dict[str, Any]) -> None:
        if node["type"] == "decision":
            rows.append((node["depth"], node["node_id"]))
            for child_node in node["children"].values():
                walk(child_node)

    walk(tree)
    rows.sort(reverse=True)
    return [node_id for _, node_id in rows]


def find_node_by_id(tree: Dict[str, Any], target_node_id: int) -> Dict[str, Any] | None:
    """Find a node inside the tree by node ID."""
    if tree["node_id"] == target_node_id:
        return tree
    for child_node in tree.get("children", {}).values():
        found = find_node_by_id(child_node, target_node_id)
        if found is not None:
            return found
    return None


def replace_node_with_leaf(tree: Dict[str, Any], target_node_id: int) -> Dict[str, Any]:
    """Return a copy of the tree with one subtree replaced by a dominant-class leaf."""
    tree_copy = copy.deepcopy(tree)
    target_node = find_node_by_id(tree_copy, target_node_id)
    if target_node is None or target_node["type"] == "leaf":
        return tree_copy

    target_node["type"] = "leaf"
    target_node["attribute"] = None
    target_node["children"] = {}
    target_node["predicted_class"] = target_node["dominant_class"]
    target_node["reason"] = "Post-pruned subtree replaced by dominant Volume class"
    return tree_copy


def apply_post_pruning(tree: Dict[str, Any], test_set: DataFrame, class_attribute: str, g_parameter: float) -> Tuple[Dict[str, Any], DataFrame]:
    """Apply the assignment post-pruning rule to the completed ID3 tree."""
    working_tree = copy.deepcopy(tree)
    q_test_count = len(test_set)
    decision_rows = []

    for node_id in collect_decision_node_ids(working_tree):
        node = find_node_by_id(working_tree, node_id)
        if node is None or node["type"] == "leaf":
            continue

        before_predictions = predict_with_id3_tree(working_tree, test_set, class_attribute)
        n_before = count_correct_predictions(before_predictions)
        k_branches = count_subtree_branches(node)
        candidate_tree = replace_node_with_leaf(working_tree, node_id)
        after_predictions = predict_with_id3_tree(candidate_tree, test_set, class_attribute)
        m_after = count_correct_predictions(after_predictions)
        left_side = (n_before - m_after) / q_test_count if q_test_count else 0
        right_side = g_parameter * k_branches
        can_prune = left_side <= right_side

        decision_rows.append({
            "Candidate Node ID": node_id,
            "Subtree Path": node.get("path_text", ""),
            "N Before Removal": n_before,
            "M After Removal": m_after,
            "Q Test Records": q_test_count,
            "K Branch Count": k_branches,
            "g Parameter": g_parameter,
            "(N-M)/Q": left_side,
            "gK": right_side,
            "Can Remove Subtree?": "Yes" if can_prune else "No",
            "Decision": "Subtree removed" if can_prune else "Subtree kept",
        })

        if can_prune:
            working_tree = candidate_tree

    return working_tree, pd.DataFrame(decision_rows)


def calculate_metrics(prediction_df: DataFrame) -> Dict[str, DataFrame]:
    """Build multiclass confusion matrix, class metrics, and overall accuracy."""
    actual_values = sorted(set(prediction_df["Actual Volume"].tolist()) | set(prediction_df["Predicted Volume"].tolist()))
    confusion_rows = []
    for actual_class in actual_values:
        row = {"Actual Volume": actual_class}
        for predicted_class in actual_values:
            row[f"Predicted {predicted_class}"] = int(
                ((prediction_df["Actual Volume"] == actual_class) & (prediction_df["Predicted Volume"] == predicted_class)).sum()
            )
        confusion_rows.append(row)
    confusion_df = pd.DataFrame(confusion_rows)

    total_records = len(prediction_df)
    correct_records = int((prediction_df["Actual Volume"] == prediction_df["Predicted Volume"]).sum())
    overall_df = pd.DataFrame([
        {
            "Total Test Records": total_records,
            "Correct Predictions": correct_records,
            "Incorrect Predictions": total_records - correct_records,
            "Accuracy": correct_records / total_records if total_records else 0,
            "Accuracy Percent": (correct_records / total_records * 100) if total_records else 0,
        }
    ])

    metric_rows = []
    for class_value in actual_values:
        tp = int(((prediction_df["Actual Volume"] == class_value) & (prediction_df["Predicted Volume"] == class_value)).sum())
        fn = int(((prediction_df["Actual Volume"] == class_value) & (prediction_df["Predicted Volume"] != class_value)).sum())
        fp = int(((prediction_df["Actual Volume"] != class_value) & (prediction_df["Predicted Volume"] == class_value)).sum())
        precision = tp / (tp + fp) if (tp + fp) else 0
        recall = tp / (tp + fn) if (tp + fn) else 0
        metric_rows.append({
            "Volume Class": class_value,
            "TP": tp,
            "FN": fn,
            "FP": fp,
            "Precision": precision,
            "Recall": recall,
        })
    class_metrics_df = pd.DataFrame(metric_rows)
    return {"Confusion Matrix": confusion_df, "Overall Metrics": overall_df, "Class Metrics": class_metrics_df}


def identify_bayes_attribute_types(training_set: DataFrame) -> Tuple[List[str], List[str], DataFrame]:
    """Separate Naive Bayes attributes into categorical and continuous groups."""
    class_attribute = CONFIG["class_attribute"]
    continuous_attributes = [attribute for attribute in CONFIG["bayes_continuous_attributes"] if attribute in training_set.columns]
    categorical_attributes = [
        column for column in training_set.columns
        if column != class_attribute and column not in continuous_attributes
    ]
    rows = []
    for attribute in categorical_attributes:
        rows.append({"Attribute": attribute, "Type Used in Bayes": "Categorical", "Reason": "Discrete product descriptor"})
    for attribute in continuous_attributes:
        rows.append({"Attribute": attribute, "Type Used in Bayes": "Continuous", "Reason": "Assignment states this attribute is not discretized"})
    return categorical_attributes, continuous_attributes, pd.DataFrame(rows)


def calculate_class_priors(training_set: DataFrame, class_attribute: str) -> DataFrame:
    """Calculate P(Volume class) for every class in the training set."""
    total_records = len(training_set)
    rows = []
    for class_value, count in sorted(Counter(training_set[class_attribute]).items()):
        rows.append({
            "Volume Class": class_value,
            "Class Count": count,
            "Total Training Records": total_records,
            "Prior Probability": count / total_records if total_records else 0,
            "Formula": f"{count} / {total_records}",
        })
    return pd.DataFrame(rows)


def calculate_categorical_counts(
    training_set: DataFrame,
    test_set: DataFrame,
    categorical_attributes: List[str],
    class_attribute: str,
) -> Tuple[DataFrame, Dict[str, List[Any]]]:
    """Count categorical attribute values by Volume class."""
    class_values = sorted(training_set[class_attribute].unique().tolist())
    attribute_domains = {
        attribute: sorted(set(training_set[attribute].tolist()) | set(test_set[attribute].tolist()))
        for attribute in categorical_attributes
    }
    rows = []
    for attribute in categorical_attributes:
        for attribute_value in attribute_domains[attribute]:
            for class_value in class_values:
                class_subset = training_set.loc[training_set[class_attribute] == class_value]
                count = int((class_subset[attribute] == attribute_value).sum())
                rows.append({
                    "Attribute": attribute,
                    "Attribute Value": attribute_value,
                    "Volume Class": class_value,
                    "Count Within Class": count,
                    "Class Count": len(class_subset),
                })
    return pd.DataFrame(rows), attribute_domains



def check_smoothing_needed(categorical_counts_df: DataFrame) -> Tuple[bool, DataFrame]:
    """Check whether any categorical count is zero, which would zero out a Bayes score."""
    zero_rows = categorical_counts_df.loc[categorical_counts_df["Count Within Class"] == 0].copy()
    smoothing_needed = not zero_rows.empty
    if smoothing_needed:
        zero_rows["Smoothing Needed?"] = "Yes"
        zero_rows["Reason"] = "A zero count would make the full Naive Bayes class score become zero."
    else:
        zero_rows = pd.DataFrame([{
            "Smoothing Needed?": "No",
            "Reason": "No zero categorical count was found across the checked attribute/class combinations.",
        }])
    return smoothing_needed, zero_rows


def calculate_categorical_probabilities(
    categorical_counts_df: DataFrame,
    attribute_domains: Dict[str, List[Any]],
    smoothing_needed: bool,
) -> DataFrame:
    """Calculate P(attribute value | Volume class) from categorical counts."""
    rows = []
    for _, row in categorical_counts_df.iterrows():
        attribute = row["Attribute"]
        count = int(row["Count Within Class"])
        class_count = int(row["Class Count"])
        domain_size = len(attribute_domains[attribute])
        if smoothing_needed:
            probability = (count + 1) / (class_count + domain_size)
            formula = f"({count} + 1) / ({class_count} + {domain_size})"
            method = "Laplace smoothing"
        else:
            probability = count / class_count if class_count else 0
            formula = f"{count} / {class_count}"
            method = "No smoothing"
        rows.append({
            "Attribute": attribute,
            "Attribute Value": row["Attribute Value"],
            "Volume Class": row["Volume Class"],
            "Count Within Class": count,
            "Class Count": class_count,
            "Domain Size": domain_size,
            "Probability": probability,
            "Formula": formula,
            "Method": method,
        })
    return pd.DataFrame(rows)


def calculate_mean(values: List[float]) -> float:
    """Calculate the average from scratch for defense-ready continuous statistics."""
    return sum(values) / len(values) if values else 0.0


def calculate_sample_std(values: List[float], mean_value: float) -> float:
    """Calculate sample standard deviation from scratch using n-1 in the denominator."""
    if len(values) < 2:
        return 0.0
    squared_total = sum((value - mean_value) ** 2 for value in values)
    variance = squared_total / (len(values) - 1)
    return math.sqrt(variance)


def calculate_continuous_statistics(
    training_set: DataFrame,
    continuous_attributes: List[str],
    class_attribute: str,
) -> DataFrame:
    """Calculate mean and sample standard deviation by Volume class for continuous attributes."""
    rows = []
    for class_value in sorted(training_set[class_attribute].unique().tolist()):
        class_subset = training_set.loc[training_set[class_attribute] == class_value]
        for attribute in continuous_attributes:
            values = [float(value) for value in class_subset[attribute].tolist()]
            mean_value = calculate_mean(values)
            std_value = calculate_sample_std(values, mean_value)
            rows.append({
                "Volume Class": class_value,
                "Attribute": attribute,
                "Record Count": len(values),
                "Mean": mean_value,
                "Sample STD": std_value,
                "STD Formula Note": "sqrt(sum((x-mean)^2)/(n-1))",
            })
    return pd.DataFrame(rows)


def calculate_continuous_probability(value: float, mean_value: float, std_value: float) -> float:
    """Calculate the Gaussian probability density component for one continuous value."""
    safe_std = std_value if std_value > 0 else 0.000000001
    exponent = -((value - mean_value) ** 2) / (2 * (safe_std ** 2))
    return (1 / (math.sqrt(2 * math.pi) * safe_std)) * math.exp(exponent)


def lookup_categorical_probability(
    probability_df: DataFrame,
    attribute: str,
    attribute_value: Any,
    class_value: Any,
) -> float:
    """Look up a categorical probability from the probability table."""
    match = probability_df.loc[
        (probability_df["Attribute"] == attribute)
        & (probability_df["Attribute Value"] == attribute_value)
        & (probability_df["Volume Class"] == class_value)
    ]
    if match.empty:
        return 0.0
    return float(match.iloc[0]["Probability"])


def lookup_continuous_stats(
    continuous_stats_df: DataFrame,
    attribute: str,
    class_value: Any,
) -> Tuple[float, float]:
    """Look up mean and standard deviation for a continuous attribute and class."""
    match = continuous_stats_df.loc[
        (continuous_stats_df["Attribute"] == attribute)
        & (continuous_stats_df["Volume Class"] == class_value)
    ]
    if match.empty:
        return 0.0, 0.0
    return float(match.iloc[0]["Mean"]), float(match.iloc[0]["Sample STD"])


def predict_bayes_test_records(
    test_set: DataFrame,
    class_priors_df: DataFrame,
    categorical_probabilities_df: DataFrame,
    continuous_stats_df: DataFrame,
    categorical_attributes: List[str],
    continuous_attributes: List[str],
    class_attribute: str,
) -> Tuple[DataFrame, DataFrame, DataFrame]:
    """Predict test records with Naive Bayes using from-scratch probability calculations."""
    class_values = sorted(class_priors_df["Volume Class"].tolist())
    prior_lookup = dict(zip(class_priors_df["Volume Class"], class_priors_df["Prior Probability"]))
    detail_rows = []
    score_rows = []
    prediction_rows = []

    for row_number, (_, row) in enumerate(test_set.iterrows(), start=1):
        best_class = None
        best_score = -1.0
        best_score_text = ""

        for class_value in class_values:
            prior_probability = float(prior_lookup[class_value])
            categorical_product = 1.0
            continuous_product = 1.0
            formula_pieces = [f"P(Volume={class_value})={prior_probability:.10f}"]

            for attribute in categorical_attributes:
                attribute_value = row[attribute]
                probability = lookup_categorical_probability(
                    categorical_probabilities_df,
                    attribute,
                    attribute_value,
                    class_value,
                )
                categorical_product *= probability
                formula_pieces.append(f"P({attribute}={attribute_value}|Volume={class_value})={probability:.10f}")
                detail_rows.append({
                    "Test Row Number": row_number,
                    "Volume Class Being Scored": class_value,
                    "Component Type": "Categorical",
                    "Attribute": attribute,
                    "Test Value": attribute_value,
                    "Probability Component": probability,
                })

            for attribute in continuous_attributes:
                attribute_value = float(row[attribute])
                mean_value, std_value = lookup_continuous_stats(continuous_stats_df, attribute, class_value)
                probability = calculate_continuous_probability(attribute_value, mean_value, std_value)
                continuous_product *= probability
                formula_pieces.append(f"f({attribute}={attribute_value}|Volume={class_value})={probability:.10f}")
                detail_rows.append({
                    "Test Row Number": row_number,
                    "Volume Class Being Scored": class_value,
                    "Component Type": "Continuous",
                    "Attribute": attribute,
                    "Test Value": attribute_value,
                    "Mean Used": mean_value,
                    "STD Used": std_value,
                    "Probability Component": probability,
                })

            final_score = prior_probability * categorical_product * continuous_product
            score_text = " * ".join(formula_pieces)
            score_rows.append({
                "Test Row Number": row_number,
                "Actual Volume": row[class_attribute],
                "Volume Class Being Scored": class_value,
                "Prior Probability": prior_probability,
                "Categorical Product": categorical_product,
                "Continuous Product": continuous_product,
                "Final Class Score": final_score,
                "Score Formula Pieces": score_text,
            })

            if final_score > best_score:
                best_score = final_score
                best_class = class_value
                best_score_text = score_text

        prediction_rows.append({
            "Test Row Number": row_number,
            "Actual Volume": row[class_attribute],
            "Predicted Volume": best_class,
            "Correct?": "Yes" if row[class_attribute] == best_class else "No",
            "Winning Score": best_score,
            "Winning Score Details": best_score_text,
            **{column: row[column] for column in test_set.columns},
        })

    return pd.DataFrame(detail_rows), pd.DataFrame(score_rows), pd.DataFrame(prediction_rows)


def build_html_table(headers: List[str], rows: List[List[Any]]) -> str:
    """Build a simple HTML table for the report."""
    html_lines = ["<table>", "<thead>", "<tr>"]
    for header in headers:
        html_lines.append(f"<th>{escape(str(header))}</th>")
    html_lines.extend(["</tr>", "</thead>", "<tbody>"])
    for row in rows:
        html_lines.append("<tr>")
        for cell in row:
            html_lines.append(f"<td>{escape(str(cell))}</td>")
        html_lines.append("</tr>")
    html_lines.extend(["</tbody>", "</table>"])
    return "\n".join(html_lines)


def build_report_text(
    method_label: str,
    user_inputs: Dict[str, Any],
    summary_tables: Dict[str, DataFrame],
    result_text: str,
) -> str:
    """Build the method-specific report text used for TXT, MD, HTML, DOCX, and PDF exports."""
    lines = []
    lines.append(f"# Assignment 2 Data Mining Report - {method_label}")
    lines.append("")
    lines.append("## Program Inputs")
    for key, value in user_inputs.items():
        lines.append(f"- **{key}:** {value}")
    lines.append("")
    lines.append("## What This Method Does")
    if method_label == "ID3":
        lines.append("ID3 builds a decision tree from the training set, extracts readable IF-THEN rules, and uses those rules to classify the test set.")
        lines.append("The pre-pruning threshold T1 controls early stopping during tree construction.  The post-pruning parameter g controls whether a completed subtree can be removed after testing the effect on classification quality.")
    else:
        lines.append("Naive Bayes calculates one probability score for each possible Volume class and predicts the class with the largest score.")
        lines.append("Categorical attributes use conditional probability tables.  Weight and Material are treated as continuous attributes and use the continuous probability formula from class notes.")
    lines.append("")
    lines.append("## Main Result")
    lines.append(result_text)
    lines.append("")
    lines.append("## Formula Notes")
    lines.append("- ID3 entropy/MC: -sum(p * log2(p)).")
    lines.append("- ID3 Gain: MC(D) - WMC(attribute).")
    lines.append("- ID3 pre-pruning mixture ratio: alpha = c1 / |F1|.")
    lines.append("- ID3 post-pruning: if (N - M) / Q <= gK, then the subtree can be removed.")
    lines.append("- Naive Bayes prior: count(Volume class) / total training records.")
    lines.append("- Naive Bayes categorical probability: count(attribute value and class) / count(class), with Laplace smoothing if needed.")
    lines.append("- Naive Bayes continuous probability uses the Gaussian density formula with mean and sample standard deviation by class.")
    lines.append("")
    lines.append("## Output Tables Created")
    for table_name in summary_tables:
        lines.append(f"- {table_name}")
    lines.append("")
    return "\n".join(lines)


def build_html_report(title: str, markdown_text: str) -> str:
    """Build a printable HTML report from the report text."""
    body_lines = []
    for line in markdown_text.splitlines():
        if line.startswith("# "):
            body_lines.append(f"<h1>{escape(line[2:])}</h1>")
        elif line.startswith("## "):
            body_lines.append(f"<h2>{escape(line[3:])}</h2>")
        elif line.startswith("- "):
            body_lines.append(f"<p>{escape(line)}</p>")
        elif line.strip() == "":
            body_lines.append("<br>")
        else:
            body_lines.append(f"<p>{escape(line)}</p>")
    return "\n".join([
        "<!DOCTYPE html>",
        "<html>",
        "<head>",
        "<meta charset='utf-8'>",
        f"<title>{escape(title)}</title>",
        "<style>",
        "body { font-family: Georgia, 'Times New Roman', serif; margin: 36px; line-height: 1.5; color: #222; }",
        "h1 { color: #7a2f20; border-bottom: 3px solid #d9b7a8; padding-bottom: 8px; }",
        "h2 { background: #f5e7df; padding: 8px; border-left: 5px solid #b85742; }",
        "p { margin: 6px 0; }",
        "</style>",
        "</head>",
        "<body>",
        *body_lines,
        "</body>",
        "</html>",
    ])


def save_docx_report(file_path: Path, title: str, markdown_text: str) -> None:
    """Save a DOCX version of the report using python-docx."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx is required to create DOCX files. Run: pip install python-docx") from exc

    document = Document()
    document.add_heading(title, level=1)
    for line in markdown_text.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip() == "":
            document.add_paragraph("")
        else:
            document.add_paragraph(line)
    document.save(file_path)
    print_saved_file(file_path)


def save_pdf_report(file_path: Path, title: str, markdown_text: str) -> None:
    """Save a simple PDF version of the report using reportlab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise ImportError("reportlab is required to create PDF files. Run: pip install reportlab") from exc

    pdf = canvas.Canvas(str(file_path), pagesize=letter)
    width, height = letter
    left_margin = 54
    y_position = height - 54
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(left_margin, y_position, title[:90])
    y_position -= 26
    pdf.setFont("Helvetica", 9)

    for raw_line in markdown_text.replace("**", "").splitlines():
        line = raw_line.replace("#", "").strip()
        if not line:
            y_position -= 8
            continue
        wrapped_lines = []
        while len(line) > CONFIG["pdf_line_width"]:
            split_at = line.rfind(" ", 0, CONFIG["pdf_line_width"])
            if split_at == -1:
                split_at = CONFIG["pdf_line_width"]
            wrapped_lines.append(line[:split_at])
            line = line[split_at:].strip()
        wrapped_lines.append(line)
        for wrapped_line in wrapped_lines:
            if y_position < 54:
                pdf.showPage()
                pdf.setFont("Helvetica", 9)
                y_position = height - 54
            pdf.drawString(left_margin, y_position, wrapped_line)
            y_position -= 12
    pdf.save()
    print_saved_file(file_path)


def save_report_files(report_dir: Path, base_name: str, title: str, report_text: str) -> None:
    """Save report text in TXT, MD, HTML, DOCX, and PDF formats."""
    save_text_output(report_dir / f"{base_name}.txt", report_text)
    save_text_output(report_dir / f"{base_name}.md", report_text)
    save_text_output(report_dir / f"{base_name}.html", build_html_report(title, report_text))
    save_docx_report(report_dir / f"{base_name}.docx", title, report_text)
    save_pdf_report(report_dir / f"{base_name}.pdf", title, report_text)


def build_defense_notes(method_label: str, user_inputs: Dict[str, Any], smoothing_text: str = "") -> str:
    """Create quick-glance notes for the live Zoom defense."""
    lines = []
    lines.append(f"# Assignment 2 Defense Notes - {method_label}")
    lines.append("")
    lines.append("## Quick Setup")
    for key, value in user_inputs.items():
        lines.append(f"- {key}: {value}")
    lines.append("")
    lines.append("## Who, What, When, Where, Why, How")
    if method_label == "ID3":
        lines.append("- Who: DFCA wants a classifier for product order Volume.")
        lines.append("- What: ID3 creates a decision tree and IF-THEN rules.")
        lines.append("- When: The tree is built from the ID3 training set and tested against the ID3 test set.")
        lines.append("- Where: The logic runs in the Python pipeline; each calculation table is exported to Excel.")
        lines.append("- Why: ID3 gives readable rules, which helps explain why a Volume class was predicted.")
        lines.append("- How: The program calculates entropy, weighted entropy, information gain, pre-pruning alpha, post-pruning loss, and final predictions.")
        lines.append("- T1: User-chosen pre-pruning threshold. Higher T1 usually stops more branches early. Lower T1 usually allows deeper splitting.")
        lines.append("- g: User-chosen post-pruning parameter, valid range 0 < g <= 0.015. Higher g allows more pruning. Lower g allows less pruning.")
    else:
        lines.append("- Who: DFCA wants a classifier for product order Volume.")
        lines.append("- What: Naive Bayes calculates probability scores for each possible Volume class.")
        lines.append("- When: Probability tables are learned from the Bayes training set, then applied to the Bayes test set.")
        lines.append("- Where: Categorical probabilities, continuous statistics, scores, and predictions are exported to Excel.")
        lines.append("- Why: Naive Bayes can handle both categorical attributes and continuous Weight/Material values.")
        lines.append("- How: The program multiplies the prior probability, categorical probability components, and continuous probability components for each class.")
        lines.append(f"- Smoothing: {smoothing_text}")
    lines.append("")
    lines.append("## Formula Checklist")
    lines.append("- Entropy/MC = -sum(p * log2(p)).")
    lines.append("- Gain = MC(D) - WMC(attribute).")
    lines.append("- ID3 alpha = c1 / |F1|.")
    lines.append("- ID3 post-pruning: if (N - M) / Q <= gK, remove subtree.")
    lines.append("- Naive Bayes prior = count(class) / total records.")
    lines.append("- Laplace smoothing = (count + 1) / (class count + number of possible values).")
    lines.append("- Continuous probability = 1/(sqrt(2*pi)*std) * e^-((value-mean)^2/(2*std^2)).")
    lines.append("")
    return "\n".join(lines)


def run_id3_pipeline(
    training_set: DataFrame,
    test_set: DataFrame,
    folders: Dict[str, Path],
    user_inputs: Dict[str, Any],
    live_output: bool,
) -> None:
    """Run the full ID3 path from dataset audit through report creation."""
    class_attribute = CONFIG["class_attribute"]
    output_dir = folders["id3_outputs"]
    report_dir = folders["id3_reports"]
    private_notes_dir = folders["id3_private_notes"]
    candidate_attributes = [column for column in training_set.columns if column != class_attribute]

    live_print(live_output, "\n--- Loading ID3 training and test files ---")
    save_excel_output(output_dir / "03_id3_training_dataset_overview.xlsx", inspect_dataset(training_set, "ID3 Training Set"))
    save_excel_output(output_dir / "04_id3_test_dataset_overview.xlsx", inspect_dataset(test_set, "ID3 Test Set"))

    live_print(live_output, "\n--- Calculating Volume class distribution ---")
    class_distribution_df = calculate_class_distribution(training_set, class_attribute)
    save_excel_output(output_dir / "05_id3_class_distribution.xlsx", class_distribution_df)

    live_print(live_output, "\n--- Building ID3 tree with pre-pruning ---")
    starting_entropy_df = pd.DataFrame([{
        "Class Attribute": class_attribute,
        "Starting MC/Entropy": calculate_entropy(training_set[class_attribute].tolist()),
        "Training Record Count": len(training_set),
        "Class Counts": format_class_counts(training_set[class_attribute].tolist()),
    }])
    save_excel_output(output_dir / "06_id3_entropy_start.xlsx", starting_entropy_df)

    gain_rows: List[Dict[str, Any]] = []
    split_rows: List[Dict[str, Any]] = []
    pre_pruning_rows: List[Dict[str, Any]] = []
    tree = build_id3_tree(
        training_set,
        candidate_attributes,
        class_attribute,
        float(user_inputs["T1 Threshold"]),
        [0],
        gain_rows,
        split_rows,
        pre_pruning_rows,
    )

    gain_df = pd.DataFrame(gain_rows)
    split_df = pd.DataFrame(split_rows)
    pre_pruning_df = pd.DataFrame(pre_pruning_rows)
    tree_df = flatten_tree(tree)
    rules_before_df = extract_id3_rules(tree)
    live_print(live_output, "\n--- Saving ID3 tree, gain, and rule outputs ---")
    save_excel_output(output_dir / "07_id3_gain_calculations_by_node.xlsx", gain_df)
    save_excel_output(output_dir / "08_id3_selected_splits.xlsx", split_df)
    save_excel_output(output_dir / "09_id3_pre_pruning_decisions.xlsx", pre_pruning_df)
    save_excel_output(output_dir / "10_id3_tree_structure_before_post_pruning.xlsx", tree_df)
    save_excel_output(output_dir / "11_id3_extracted_rules_before_post_pruning.xlsx", rules_before_df)

    live_print(live_output, "\n--- Testing ID3 predictions before post-pruning ---")
    predictions_before_df = predict_with_id3_tree(tree, test_set, class_attribute)
    metrics_before = calculate_metrics(predictions_before_df)
    save_excel_output(output_dir / "12_id3_test_predictions_before_post_pruning.xlsx", predictions_before_df)
    save_excel_output(output_dir / "13_id3_metrics_before_post_pruning.xlsx", metrics_before)

    live_print(live_output, "\n--- Applying ID3 post-pruning ---")
    pruned_tree, post_pruning_df = apply_post_pruning(tree, test_set, class_attribute, float(user_inputs["g Parameter"] ))
    tree_after_df = flatten_tree(pruned_tree)
    rules_after_df = extract_id3_rules(pruned_tree)
    predictions_after_df = predict_with_id3_tree(pruned_tree, test_set, class_attribute)
    metrics_after = calculate_metrics(predictions_after_df)
    live_print(live_output, "\n--- Saving post-pruning tree, rule, and metric outputs ---")
    save_excel_output(output_dir / "14_id3_post_pruning_calculations.xlsx", post_pruning_df)
    save_excel_output(output_dir / "15_id3_tree_structure_after_post_pruning.xlsx", tree_after_df)
    save_excel_output(output_dir / "16_id3_extracted_rules_after_post_pruning.xlsx", rules_after_df)
    save_excel_output(output_dir / "17_id3_test_predictions_after_post_pruning.xlsx", predictions_after_df)
    save_excel_output(output_dir / "18_id3_metrics_after_post_pruning.xlsx", metrics_after)

    threshold_explanations_df = pd.DataFrame([
        {
            "Threshold/Parameter": "T1",
            "Used In": "ID3 pre-pruning",
            "User Chosen?": "Yes",
            "Valid Range": "0 < T1 <= 1",
            "Meaning": "Controls when a branch stops splitting early based on alpha = c1 / |F1|.",
            "If Increased": "More branches may stop early, creating a smaller tree.",
            "If Decreased": "More branches may continue splitting, creating a deeper tree.",
        },
        {
            "Threshold/Parameter": "g",
            "Used In": "ID3 post-pruning",
            "User Chosen?": "Yes",
            "Valid Range": "0 < g <= 0.015",
            "Meaning": "Controls whether accuracy loss is small enough to remove a subtree.",
            "If Increased": "More subtrees may be removed.",
            "If Decreased": "Fewer subtrees may be removed.",
        },
    ])
    save_excel_output(output_dir / "19_id3_threshold_explanations.xlsx", threshold_explanations_df)

    final_accuracy = float(metrics_after["Overall Metrics"].iloc[0]["Accuracy Percent"])
    result_text = f"After post-pruning, ID3 correctly classified {final_accuracy:.2f}% of the ID3 test records."
    summary_tables = {
        "ID3 Class Distribution": class_distribution_df,
        "ID3 Gain Calculations": gain_df,
        "ID3 Pre-Pruning Decisions": pre_pruning_df,
        "ID3 Post-Pruning Decisions": post_pruning_df,
        "ID3 Final Metrics": metrics_after["Overall Metrics"],
    }
    report_text = build_report_text("ID3", user_inputs, summary_tables, result_text)
    defense_notes = build_defense_notes("ID3", user_inputs)
    live_print(live_output, "\n--- Saving final ID3 report package ---")
    save_report_files(report_dir, "Assignment2_Report_ID3", "Assignment 2 Report - ID3", report_text)
    save_excel_output(private_notes_dir / "20_id3_defense_notes.xlsx", pd.DataFrame({"Defense Notes": defense_notes.splitlines()}))
    save_report_files(private_notes_dir, "Assignment2_DefenseNotes_ID3", "Assignment 2 Defense Notes - ID3", defense_notes)

    run_summary_df = pd.DataFrame([{
        "Classifier": "ID3",
        "Training Records": len(training_set),
        "Test Records": len(test_set),
        "T1 Threshold": user_inputs["T1 Threshold"],
        "g Parameter": user_inputs["g Parameter"],
        "Rules Before Post-Pruning": len(rules_before_df),
        "Rules After Post-Pruning": len(rules_after_df),
        "Final Accuracy Percent": final_accuracy,
        "Run Completed At": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }])
    live_print(live_output, "\n--- Saving ID3 run summary files ---")
    save_excel_output(output_dir / "00_run_summary.xlsx", run_summary_df)
    save_text_output(output_dir / "00_run_summary.json", run_summary_df.iloc[0].to_json(indent=2))


def run_bayes_pipeline(
    training_set: DataFrame,
    test_set: DataFrame,
    folders: Dict[str, Path],
    user_inputs: Dict[str, Any],
    live_output: bool,
) -> None:
    """Run the full Naive Bayes path from dataset audit through report creation."""
    class_attribute = CONFIG["class_attribute"]
    output_dir = folders["bayes_outputs"]
    report_dir = folders["bayes_reports"]
    private_notes_dir = folders["bayes_private_notes"]

    live_print(live_output, "\n--- Loading Bayes training and test files ---")
    save_excel_output(output_dir / "03_bayes_training_dataset_overview.xlsx", inspect_dataset(training_set, "Bayes Training Set"))
    save_excel_output(output_dir / "04_bayes_test_dataset_overview.xlsx", inspect_dataset(test_set, "Bayes Test Set"))

    live_print(live_output, "\n--- Calculating Bayes class distribution and attribute types ---")
    class_distribution_df = calculate_class_distribution(training_set, class_attribute)
    categorical_attributes, continuous_attributes, attribute_types_df = identify_bayes_attribute_types(training_set)
    save_excel_output(output_dir / "05_bayes_class_distribution.xlsx", class_distribution_df)
    save_excel_output(output_dir / "06_bayes_attribute_types.xlsx", attribute_types_df)

    live_print(live_output, "\n--- Calculating Naive Bayes probabilities from scratch ---")
    class_priors_df = calculate_class_priors(training_set, class_attribute)
    categorical_counts_df, attribute_domains = calculate_categorical_counts(training_set, test_set, categorical_attributes, class_attribute)
    smoothing_needed, smoothing_df = check_smoothing_needed(categorical_counts_df)
    categorical_probabilities_df = calculate_categorical_probabilities(categorical_counts_df, attribute_domains, smoothing_needed)
    continuous_stats_df = calculate_continuous_statistics(training_set, continuous_attributes, class_attribute)
    live_print(live_output, "\n--- Saving Bayes probability setup outputs ---")
    save_excel_output(output_dir / "07_bayes_class_priors.xlsx", class_priors_df)
    save_excel_output(output_dir / "08_bayes_categorical_counts.xlsx", categorical_counts_df)
    save_excel_output(output_dir / "09_bayes_smoothing_decision.xlsx", smoothing_df)
    save_excel_output(output_dir / "10_bayes_categorical_probabilities.xlsx", categorical_probabilities_df)
    save_excel_output(output_dir / "11_bayes_continuous_statistics.xlsx", continuous_stats_df)

    live_print(live_output, "\n--- Predicting Bayes test records ---")
    probability_details_df, class_scores_df, predictions_df = predict_bayes_test_records(
        test_set,
        class_priors_df,
        categorical_probabilities_df,
        continuous_stats_df,
        categorical_attributes,
        continuous_attributes,
        class_attribute,
    )
    metrics = calculate_metrics(predictions_df)
    live_print(live_output, "\n--- Saving Bayes prediction and metric outputs ---")
    save_excel_output(output_dir / "12_bayes_test_probability_details.xlsx", probability_details_df)
    save_excel_output(output_dir / "13_bayes_class_scores.xlsx", class_scores_df)
    save_excel_output(output_dir / "14_bayes_test_predictions.xlsx", predictions_df)
    save_excel_output(output_dir / "15_bayes_confusion_matrix.xlsx", metrics["Confusion Matrix"])
    save_excel_output(output_dir / "16_bayes_metrics.xlsx", {"Overall Metrics": metrics["Overall Metrics"], "Class Metrics": metrics["Class Metrics"]})

    smoothing_text = (
        "Laplace smoothing was used because at least one categorical value/class combination had a zero count."
        if smoothing_needed
        else "Smoothing was not used because no zero categorical probability was detected."
    )
    final_accuracy = float(metrics["Overall Metrics"].iloc[0]["Accuracy Percent"])
    result_text = f"Naive Bayes correctly classified {final_accuracy:.2f}% of the Bayes test records. {smoothing_text}"
    summary_tables = {
        "Bayes Class Distribution": class_distribution_df,
        "Bayes Class Priors": class_priors_df,
        "Bayes Smoothing Decision": smoothing_df,
        "Bayes Continuous Statistics": continuous_stats_df,
        "Bayes Final Metrics": metrics["Overall Metrics"],
    }
    report_text = build_report_text("Naive Bayes", user_inputs, summary_tables, result_text)
    defense_notes = build_defense_notes("Naive Bayes", user_inputs, smoothing_text)
    live_print(live_output, "\n--- Saving final Bayes report package ---")
    save_report_files(report_dir, "Assignment2_Report_Bayes", "Assignment 2 Report - Naive Bayes", report_text)
    save_excel_output(private_notes_dir / "17_bayes_defense_notes.xlsx", pd.DataFrame({"Defense Notes": defense_notes.splitlines()}))
    save_report_files(private_notes_dir, "Assignment2_DefenseNotes_Bayes", "Assignment 2 Defense Notes - Naive Bayes", defense_notes)

    run_summary_df = pd.DataFrame([{
        "Classifier": "Naive Bayes",
        "Training Records": len(training_set),
        "Test Records": len(test_set),
        "Categorical Attributes": ", ".join(categorical_attributes),
        "Continuous Attributes": ", ".join(continuous_attributes),
        "Smoothing Used": "Yes" if smoothing_needed else "No",
        "Final Accuracy Percent": final_accuracy,
        "Run Completed At": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }])
    live_print(live_output, "\n--- Saving Bayes run summary files ---")
    save_excel_output(output_dir / "00_run_summary.xlsx", run_summary_df)
    save_text_output(output_dir / "00_run_summary.json", run_summary_df.iloc[0].to_json(indent=2))


def main() -> None:
    """Run Assignment 2 from user input through selected classifier outputs."""
    SESSION_FLAGS["show_live_terminal_output"] = True

    classifier_choice, live_output, user_inputs = collect_startup_inputs()
    SESSION_FLAGS["show_live_terminal_output"] = live_output

    validate_runtime_environment()
    base_dir = CONFIG["base_dir"]
    validate_project_files(base_dir)
    folders = create_assignment_folders(base_dir)

    live_print(live_output, "\n--- Creating output folders ---")
    function_purpose_df = pd.DataFrame(list(explain_function_purpose().items()), columns=["Function", "Purpose"])
    method_folder = folders["id3_outputs"] if classifier_choice == "id3" else folders["bayes_outputs"]
    save_excel_output(method_folder / "01_function_purposes.xlsx", function_purpose_df)

    training_set, test_set, train_file, test_file = load_selected_datasets(classifier_choice, folders)
    user_inputs["Training File"] = train_file
    user_inputs["Test File"] = test_file
    save_excel_output(method_folder / "02_user_inputs.xlsx", pd.DataFrame([user_inputs]))

    if classifier_choice == "id3":
        run_id3_pipeline(training_set, test_set, folders, user_inputs, live_output)
    else:
        run_bayes_pipeline(training_set, test_set, folders, user_inputs, live_output)

    print("\nRun complete.", flush=True)


if __name__ == "__main__":
    try:
        main()
    except (ImportError, RuntimeError, FileNotFoundError) as exc:
        print(f"\nStartup check failed: {exc}", flush=True)
        raise SystemExit(1)
